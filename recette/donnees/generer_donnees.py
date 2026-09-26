#!/usr/bin/env python3
"""Génère les jeux de données de recette de la GED Marchica Med.

Pourquoi un générateur plutôt que des fichiers « trouvés » : chaque fichier porte
un mot-témoin connu (pour prouver qu'il est retrouvé par son contenu en E6), son
type réel est maîtrisé (pour les contrôles Tika en E5) et il ne contient aucune
donnée réelle de Marchica Med.

Dépendances : PyMuPDF (fitz), reportlab, Pillow, numpy, python-docx.
Polices : une police couvrant l'arabe est nécessaire pour les scans arabes
(Arial ou Traditional Arabic sous Windows ; sous Linux, Noto Naskh Arabic
ou DejaVu Sans) : voir --police-arabe.

Usage :
    python generer_donnees.py                    # jeu versionné (dossier courant du script)
    python generer_donnees.py --pages-scan 20 --sortie genere/
                                                 # scans de 20 pages (critère de sortie E6),
                                                 # trop volumineux pour être versionnés
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import os
import sys
from pathlib import Path

import fitz  # PyMuPDF
import numpy as np
from PIL import Image, ImageFilter
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ICI = Path(__file__).resolve().parent

# Mots-témoins : inventés, absents de tout dictionnaire, pour qu'une recherche
# plein texte qui les retrouve prouve la lecture du contenu et non un hasard.
TEMOIN_FR = "zarkolinet"
TEMOIN_AR = "زركولين"

POLICES_ARABES = [
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/trado.ttf",
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]

TEXTE_COURRIER_FR = [
    "Agence pour l'Aménagement du Site de la Lagune de Marchica",
    "Direction Administrative et Financière — Bureau d'ordre",
    "Nador, le 26 septembre 2026",
    "Objet : convention de partenariat relative à l'aménagement des berges de la lagune",
    f"Référence de recette : {TEMOIN_FR}",
    "Monsieur le Directeur,",
    "Nous avons l'honneur de vous transmettre, pour signature, la convention de partenariat "
    "portant sur la réhabilitation écologique des berges et la création d'un sentier pédestre "
    "de sept kilomètres. Le montant prévisionnel des travaux s'élève à 4 250 000,00 dirhams "
    "toutes taxes comprises, réparti sur les exercices 2026 et 2027.",
    "Les pièces jointes comprennent le décompte provisoire numéro 14, le procès-verbal de la "
    "réunion de coordination du 12 septembre 2026 et le planning d'exécution révisé.",
    "Veuillez agréer, Monsieur le Directeur, l'expression de notre haute considération.",
]

TEXTE_COURRIER_AR = [
    "المملكة المغربية",
    "وكالة تهيئة موقع بحيرة مارتشيكا",
    "مكتب الضبط — المديرية الإدارية والمالية",
    "الناظور في 26 شتنبر 2026",
    "الموضوع : اتفاقية شراكة تتعلق بتهيئة ضفاف البحيرة",
    f"مرجع الاختبار : {TEMOIN_AR}",
    "السيد المدير المحترم،",
    "يشرفنا أن نوافيكم قصد التوقيع باتفاقية الشراكة المتعلقة بإعادة التأهيل البيئي لضفاف "
    "البحيرة وإحداث ممر للراجلين على طول سبعة كيلومترات. ويبلغ المبلغ التقديري للأشغال "
    "أربعة ملايين ومائتين وخمسين ألف درهم مع احتساب جميع الرسوم.",
    "وتتضمن الوثائق المرفقة الكشف المؤقت رقم 14 ومحضر اجتماع التنسيق المنعقد بتاريخ "
    "12 شتنبر 2026 والجدول الزمني المعدل للإنجاز.",
    "وتقبلوا، السيد المدير، فائق عبارات التقدير والاحترام.",
]


# ---------------------------------------------------------------- utilitaires

def sha256(chemin: Path) -> str:
    h = hashlib.sha256()
    with open(chemin, "rb") as f:
        for bloc in iter(lambda: f.read(1 << 16), b""):
            h.update(bloc)
    return h.hexdigest()


def police_arabe(imposee: str | None) -> Path:
    candidats = [imposee] if imposee else POLICES_ARABES
    for c in candidats:
        if c and Path(c).is_file():
            return Path(c)
    sys.exit("Aucune police couvrant l'arabe trouvée : passer --police-arabe <fichier.ttf>.")


def page_html(lignes: list[str], rtl: bool, police: Path | None, numero: int, total: int) -> fitz.Document:
    """Compose une page A4 par le moteur HTML de MuPDF.

    MuPDF met en forme l'arabe avec HarfBuzz (ligatures, formes contextuelles,
    sens RTL) : c'est ce qui permet d'obtenir un rendu fidèle sans dépendre de
    libraqm, absente des Pillow distribués sous Windows.
    """
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    direction = "rtl" if rtl else "ltr"
    align = "right" if rtl else "justify"
    css = f"* {{font-size: 12pt; line-height: 1.5;}} p {{direction: {direction}; text-align: {align}; margin: 0 0 9pt 0;}}"
    archive = None
    if police is not None:
        archive = fitz.Archive(str(police.parent))
        css = f"@font-face {{font-family: corps; src: url({police.name});}} * {{font-family: corps;}} " + css
    paragraphes = "".join(f"<p>{l}</p>" for l in lignes)
    pied = f"<p>الصفحة {numero} من {total}</p>" if rtl else f"<p>— {numero} / {total} —</p>"
    page.insert_htmlbox(fitz.Rect(60, 60, 535, 780), paragraphes + pied, css=css, archive=archive)
    return doc


def degrader(img: Image.Image, graine: int) -> Image.Image:
    """Imite un scanner de bureau : léger biais, grain, fond non blanc.

    Assez pour que la page ne soit pas un rendu vectoriel parfait, pas assez
    pour rendre l'OCR impossible : le but est un cas réaliste, pas un piège.
    """
    rng = np.random.default_rng(graine)
    img = img.convert("L").rotate(0.6 if graine % 2 else -0.5, resample=Image.BICUBIC, fillcolor=255)
    a = np.asarray(img, dtype=np.int16)
    a = a - 12  # fond légèrement gris
    taches = rng.random(a.shape)
    a[taches < 0.0015] = 40      # poussières sombres
    a[taches > 0.9985] = 255     # manques d'encre
    a = np.clip(a, 0, 255).astype(np.uint8)
    return Image.fromarray(a, "L").filter(ImageFilter.GaussianBlur(0.4))


def pdf_scanne(sortie: Path, lignes: list[str], rtl: bool, police: Path | None, pages: int, dpi: int = 300) -> None:
    """PDF composé uniquement d'images (aucune couche texte), comme un vrai scan."""
    resultat = fitz.open()
    for n in range(1, pages + 1):
        source = page_html(lignes, rtl, police, n, pages)
        pix = source[0].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        img = degrader(Image.frombytes("L", (pix.width, pix.height), pix.samples), graine=n)
        tampon = io.BytesIO()
        img.save(tampon, format="JPEG", quality=72, dpi=(dpi, dpi))
        page = resultat.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=tampon.getvalue())
    resultat.save(sortie, garbage=4, deflate=True)
    # Garde-fou : un « scan » qui porterait une couche texte fausserait le test OCR.
    verif = fitz.open(sortie)
    if any(p.get_text().strip() for p in verif):
        sys.exit(f"{sortie} contient une couche texte : ce n'est pas un scan.")


# ---------------------------------------------------------------- fichiers

def pdf_texte_reportlab(sortie: Path, titre: str, lignes: list[str]) -> None:
    # invariant : octets identiques d'une génération à l'autre (pas de date de
    # création), sinon le manifeste changerait à chaque exécution.
    c = canvas.Canvas(str(sortie), pagesize=A4, invariant=1)
    c.setTitle(titre)
    c.setAuthor("Recette GED — données fictives")
    largeur, hauteur = A4
    for num_page in (1, 2):
        y = hauteur - 70
        c.setFont("Helvetica-Bold", 13)
        c.drawString(60, y, titre)
        y -= 30
        c.setFont("Helvetica", 10.5)
        for ligne in lignes:
            # Découpe naïve à 95 caractères : suffisant pour un texte de test.
            while ligne:
                c.drawString(60, y, ligne[:95])
                ligne = ligne[95:]
                y -= 15
            y -= 6
        c.drawString(60, 50, f"Page {num_page} / 2")
        c.showPage()
    c.save()


def pdf_texte_arabe(sortie: Path, police: Path) -> None:
    doc = page_html(TEXTE_COURRIER_AR, True, police, 1, 1)
    doc.set_metadata({"title": "Courrier arabe — couche texte native", "author": "Recette GED — données fictives"})
    doc.subset_fonts()  # la police complète pèserait 600 Ko pour une page
    doc.save(sortie, garbage=4, deflate=True)


def png_scan(sortie: Path) -> None:
    source = page_html(TEXTE_COURRIER_FR[:7], False, None, 1, 1)
    pix = source[0].get_pixmap(dpi=200, colorspace=fitz.csGRAY, clip=fitz.Rect(40, 40, 555, 470))
    img = degrader(Image.frombytes("L", (pix.width, pix.height), pix.samples), graine=7)
    img.save(sortie, format="PNG", optimize=True, dpi=(200, 200))


def docx_fr(sortie: Path) -> None:
    d = Document()
    d.core_properties.author = "Recette GED — données fictives"
    d.core_properties.title = "Procès-verbal de réunion de coordination"
    d.add_heading("Procès-verbal de la réunion de coordination du 12 septembre 2026", level=1)
    d.add_paragraph(f"Référence de recette : {TEMOIN_FR}-docx")
    d.add_paragraph(
        "Présents : la Direction Technique, la Direction Administrative et Financière, "
        "le bureau d'études et l'entreprise titulaire du lot 3 (berges et sentier)."
    )
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Point", "Décision", "Échéance"
    for ligne in [("Planning", "Validé avec réserve", "30/09/2026"),
                  ("Décompte n° 14", "Transmis pour visa", "05/10/2026"),
                  ("Sentier pédestre", "Tracé définitif arrêté", "15/10/2026")]:
        cellules = t.add_row().cells
        for i, v in enumerate(ligne):
            cellules[i].text = v
    d.save(sortie)


def faux_pdf_texte(sortie: Path) -> None:
    # Extension .pdf, contenu texte brut : Tika doit détecter text/plain.
    # Refus attendu (415) pour un type documentaire qui n'accepte que PDF.
    sortie.write_text(
        "Ceci n'est pas un PDF. Fichier de recette : extension .pdf, contenu texte brut.\n"
        "Le type réel doit être déterminé par le contenu (DAT V3, 6.1.5).\n",
        encoding="utf-8",
    )


def faux_pdf_executable(sortie: Path) -> None:
    # En-tête MZ/PE minimal, non exécutable (aucun code), déguisé en .pdf.
    # Tika le reconnaît comme application/x-msdownload : refus 415 attendu.
    entete = bytearray(512)
    entete[0:2] = b"MZ"
    entete[0x3C:0x40] = (0x80).to_bytes(4, "little")   # e_lfanew -> en-tête PE
    entete[0x80:0x84] = b"PE\x00\x00"
    entete[0x84:0x86] = (0x14C).to_bytes(2, "little")  # machine i386
    entete[0x40:0x40 + 39] = b"This program cannot be run in DOS mode."
    sortie.write_bytes(bytes(entete))


def textes_autorises(dossier: Path) -> list[Path]:
    txt = dossier / "note_texte_brut.txt"
    txt.write_text(f"Note interne de recette ({TEMOIN_FR}-txt).\nFormat texte brut autorisé par défaut (6.1.5).\n",
                   encoding="utf-8")
    csv_ = dossier / "tableau_decomptes.csv"
    with open(csv_, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["numero", "date", "montant_mad", "statut"])
        w.writerows([[12, "2026-07-31", "812000.00", "payé"], [13, "2026-08-31", "905500.00", "payé"],
                     [14, "2026-09-25", "1020000.00", "en visa"]])
    return [txt, csv_]


# ---------------------------------------------------------------- manifeste

ATTENDUS = {
    "pdf_texte_fr_convention.pdf": ("application/pdf", "201 (ou 202 si OCR en attente)", f"Couche texte native ; témoin « {TEMOIN_FR} »"),
    "pdf_texte_fr_facture.pdf": ("application/pdf", "201 (ou 202)", "Couche texte native, montants et dates"),
    "pdf_texte_ar_courrier.pdf": ("application/pdf", "201 (ou 202)", f"Couche texte arabe (formes de présentation Unicode) ; témoin « {TEMOIN_AR} »"),
    "scan_fr_courrier.pdf": ("application/pdf", "202 EN_ATTENTE_OCR", f"Images seules 300 dpi, aucune couche texte ; témoin « {TEMOIN_FR} » après OCR fra"),
    "scan_ar_courrier.pdf": ("application/pdf", "202 EN_ATTENTE_OCR", f"Images seules 300 dpi, arabe ; témoin « {TEMOIN_AR} » après OCR ara"),
    "image_scan_fr.png": ("image/png", "201 (ou 202)", "Image 200 dpi, OCR fra"),
    "document_fr.docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "201", f"Témoin « {TEMOIN_FR}-docx » ; prévisualisation via LibreOffice"),
    "note_texte_brut.txt": ("text/plain", "201 si le type l'autorise", "Format autorisé par défaut"),
    "tableau_decomptes.csv": ("text/csv", "201 si le type l'autorise", "Format autorisé par défaut"),
    "faux_pdf_texte.pdf": ("text/plain", "415", "Extension .pdf, contenu texte : refus par type réel"),
    "faux_pdf_executable.pdf": ("application/x-msdownload", "415", "Extension .pdf, en-tête MZ/PE : refus par type réel"),
}


def ecrire_manifeste(dossier: Path, fichiers: list[Path]) -> None:
    with open(dossier / "MANIFESTE.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["fichier", "octets", "sha256", "type_reel_attendu", "reponse_depot_attendue", "remarque"])
        for p in sorted(fichiers, key=lambda x: x.name):
            type_reel, reponse, remarque = ATTENDUS.get(p.name, ("", "", "généré à la demande"))
            w.writerow([p.name, p.stat().st_size, sha256(p), type_reel, reponse, remarque])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sortie", type=Path, default=ICI, help="dossier de sortie (défaut : dossier du script)")
    ap.add_argument("--pages-scan", type=int, default=2, help="pages des scans FR et AR (20 pour le critère E6)")
    ap.add_argument("--police-arabe", help="fichier .ttf couvrant l'arabe")
    args = ap.parse_args()

    sortie: Path = args.sortie
    sortie.mkdir(parents=True, exist_ok=True)
    police = police_arabe(args.police_arabe)
    faits: list[Path] = []

    if args.pages_scan != 2:
        # Mode volumineux : seulement les scans, nommés par leur nombre de pages.
        for nom, lignes, rtl, pol in (("scan_fr", TEXTE_COURRIER_FR, False, None),
                                      ("scan_ar", TEXTE_COURRIER_AR, True, police)):
            p = sortie / f"{nom}_{args.pages_scan}p.pdf"
            pdf_scanne(p, lignes, rtl, pol, args.pages_scan)
            print(f"{p.name}  {p.stat().st_size} octets  sha256={sha256(p)}")
        return

    p = sortie / "pdf_texte_fr_convention.pdf"; pdf_texte_reportlab(p, "Convention de partenariat — lagune de Marchica", TEXTE_COURRIER_FR); faits.append(p)
    p = sortie / "pdf_texte_fr_facture.pdf"
    pdf_texte_reportlab(p, "Facture n° F-2026-0914", [
        "Fournisseur : Société fictive de travaux lagunaires SARL", "Date de facture : 14/09/2026",
        "Désignation : aménagement des berges, lot 3, situation n° 14",
        "Montant HT : 850 000,00 MAD — TVA 20 % : 170 000,00 MAD — Montant TTC : 1 020 000,00 MAD",
        f"Référence de recette : {TEMOIN_FR}-facture"]); faits.append(p)
    p = sortie / "pdf_texte_ar_courrier.pdf"; pdf_texte_arabe(p, police); faits.append(p)
    p = sortie / "scan_fr_courrier.pdf"; pdf_scanne(p, TEXTE_COURRIER_FR, False, None, 2); faits.append(p)
    p = sortie / "scan_ar_courrier.pdf"; pdf_scanne(p, TEXTE_COURRIER_AR, True, police, 2); faits.append(p)
    p = sortie / "image_scan_fr.png"; png_scan(p); faits.append(p)
    p = sortie / "document_fr.docx"; docx_fr(p); faits.append(p)
    faits += textes_autorises(sortie)
    p = sortie / "faux_pdf_texte.pdf"; faux_pdf_texte(p); faits.append(p)
    p = sortie / "faux_pdf_executable.pdf"; faux_pdf_executable(p); faits.append(p)

    ecrire_manifeste(sortie, faits)
    for f in faits:
        print(f"{f.name:32s} {f.stat().st_size:>9d} octets")


if __name__ == "__main__":
    main()
